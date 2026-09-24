from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from pptx import Presentation
from pptx.util import Inches, Pt


_MARKDOWN_INLINE_RE = re.compile(r"(\*\*|__|~~|\`)")
_LINK_RE = re.compile(r"\[([^\]]+)\]\([^\)]+\)")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.+?)\s*$")
_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s+(.+?)\s*$")


@dataclass(frozen=True, slots=True)
class MeetingReportExports:
    markdown: Path
    docx: Path
    pptx: Path

    def paths(self) -> tuple[Path, Path, Path]:
        return self.markdown, self.docx, self.pptx


def _plain_text(value: str) -> str:
    text = value.strip()
    while text.startswith(">"):
        text = text[1:].lstrip()
    text = _LINK_RE.sub(r"\1", text)
    text = _MARKDOWN_INLINE_RE.sub("", text)
    return text.strip()


def _iter_blocks(markdown: str):
    for raw in markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            yield ("blank", "", 0)
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            yield ("heading", _plain_text(heading.group(2)), len(heading.group(1)))
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            yield ("bullet", _plain_text(bullet.group(1)), 0)
            continue

        numbered = _NUMBER_RE.match(line)
        if numbered:
            yield ("number", _plain_text(numbered.group(1)), 0)
            continue

        if line.lstrip().startswith(">"):
            yield ("quote", _plain_text(line), 0)
            continue

        yield ("paragraph", _plain_text(line), 0)


def export_markdown_to_docx(markdown: str, output_path: Path, *, title: str) -> Path:
    document = Document()
    core = document.core_properties
    core.title = title
    core.subject = "Discord meeting report"
    core.comments = "Generated from the approved Markdown meeting report."

    blocks = list(_iter_blocks(markdown))
    first_heading_consumed = False
    for kind, text, level in blocks:
        if kind == "blank":
            continue
        if not text:
            continue

        if kind == "heading":
            if not first_heading_consumed and level == 1:
                document.add_heading(text, level=0)
                first_heading_consumed = True
            else:
                document.add_heading(text, level=min(level, 4))
            continue

        if kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
            continue

        if kind == "number":
            document.add_paragraph(text, style="List Number")
            continue

        paragraph = document.add_paragraph()
        if kind == "quote":
            paragraph.style = document.styles["Quote"]
        paragraph.add_run(text)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def _ppt_sections(markdown: str, fallback_title: str) -> tuple[str, list[tuple[str, list[str]]]]:
    deck_title = fallback_title
    sections: list[tuple[str, list[str]]] = []
    current_title = "週會摘要"
    current_lines: list[str] = []

    for kind, text, level in _iter_blocks(markdown):
        if not text:
            continue

        if kind == "heading" and level == 1:
            deck_title = text
            continue

        if kind == "heading" and level <= 3:
            if current_lines:
                sections.append((current_title, current_lines))
            current_title = text
            current_lines = []
            continue

        prefix = ""
        if kind == "number":
            prefix = "• "
        elif kind == "bullet":
            prefix = "• "
        current_lines.append(prefix + text)

    if current_lines:
        sections.append((current_title, current_lines))

    if not sections:
        sections = [("週會摘要", ["完整內容請參考 Markdown / Word 版本。"])]

    return deck_title, sections


def _chunk_lines(lines: list[str], *, max_items: int = 8, max_chars: int = 900) -> list[list[str]]:
    chunks: list[list[str]] = []
    current: list[str] = []
    current_chars = 0
    for line in lines:
        normalized = line.strip()
        if not normalized:
            continue
        projected = current_chars + len(normalized)
        if current and (len(current) >= max_items or projected > max_chars):
            chunks.append(current)
            current = []
            current_chars = 0
        current.append(normalized)
        current_chars += len(normalized)
    if current:
        chunks.append(current)
    return chunks or [["完整內容請參考 Markdown / Word 版本。"]]


def export_markdown_to_pptx(markdown: str, output_path: Path, *, title: str) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    deck_title, sections = _ppt_sections(markdown, title)

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = deck_title
    subtitle = title_slide.placeholders[1]
    subtitle.text = "Discord 週會報匯出"

    for section_title, lines in sections:
        chunks = _chunk_lines(lines)
        for index, chunk in enumerate(chunks, start=1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            suffix = f" ({index}/{len(chunks)})" if len(chunks) > 1 else ""
            slide.shapes.title.text = f"{section_title}{suffix}"
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for item_index, line in enumerate(chunk):
                paragraph = frame.paragraphs[0] if item_index == 0 else frame.add_paragraph()
                paragraph.text = line.removeprefix("• ").strip()
                paragraph.level = 0
                paragraph.font.size = Pt(22)
                if line.startswith("• "):
                    paragraph.text = line[2:].strip()
                    paragraph.level = 0

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return output_path


def build_meeting_report_exports(report_path: Path, *, title: str) -> MeetingReportExports:
    markdown = report_path.read_text(encoding="utf-8")
    if not markdown.strip():
        raise ValueError("週會報內容為空，無法匯出文件。")

    stem = report_path.stem
    docx_path = report_path.with_name(f"{stem}.docx")
    pptx_path = report_path.with_name(f"{stem}.pptx")
    export_markdown_to_docx(markdown, docx_path, title=title)
    export_markdown_to_pptx(markdown, pptx_path, title=title)
    return MeetingReportExports(markdown=report_path, docx=docx_path, pptx=pptx_path)
