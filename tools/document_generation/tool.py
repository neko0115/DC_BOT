from __future__ import annotations

import os
import re
import uuid
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt as DocxPt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


MAX_TITLE = 180
MAX_SECTION_BODY = 6000
MAX_SLIDES = 20
MAX_BULLETS = 10
MAX_BULLET_CHARS = 500
MAX_NOTES = 6000

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MARKDOWN_MIME = "text/markdown"


class DocumentGenerationError(RuntimeError):
    pass


def _project_root() -> Path:
    return Path(__file__).resolve().parents[2]


def _output_dir() -> Path:
    path = _project_root() / "data" / "tool-artifacts" / "document_generation"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _text(value: object, *, label: str, maximum: int, allow_empty: bool = False) -> str:
    result = str(value or "").strip()
    if not result and not allow_empty:
        raise ValueError(f"{label} is required.")
    if len(result) > maximum:
        raise ValueError(f"{label} exceeds {maximum} characters.")
    return result


def _safe_stem(value: object, fallback: str) -> str:
    raw = str(value or fallback).strip()
    stem = re.sub(r"[^0-9A-Za-z\u3400-\u9fff._-]+", "_", raw).strip("._-")
    return (stem[:80] or fallback)


def _artifact(path: Path, *, filename: str, mime_type: str) -> dict[str, object]:
    return {
        "relative_path": f"document_generation/{path.name}",
        "filename": filename,
        "mime_type": mime_type,
        "delete_after_send": True,
    }


def _slides(arguments: dict[str, object]) -> list[dict[str, object]]:
    raw = arguments.get("slides")
    if not isinstance(raw, list) or not raw:
        raise ValueError("slides must be a non-empty array.")
    result: list[dict[str, object]] = []
    for index, item in enumerate(raw[:MAX_SLIDES], start=1):
        if not isinstance(item, dict):
            raise ValueError(f"slides[{index}] must be an object.")
        title = _text(item.get("title"), label=f"slides[{index}].title", maximum=180)
        bullets_raw = item.get("bullets", [])
        if not isinstance(bullets_raw, list):
            raise ValueError(f"slides[{index}].bullets must be an array.")
        bullets = [
            _text(value, label=f"slides[{index}].bullets", maximum=MAX_BULLET_CHARS)
            for value in bullets_raw[:MAX_BULLETS]
            if str(value or "").strip()
        ]
        notes = _text(
            item.get("speaker_notes", ""),
            label=f"slides[{index}].speaker_notes",
            maximum=MAX_NOTES,
            allow_empty=True,
        )
        result.append({"title": title, "bullets": bullets, "speaker_notes": notes})
    return result


def _theme(theme: str) -> dict[str, RGBColor]:
    if theme == "academic":
        return {
            "bg": RGBColor(245, 246, 248),
            "panel": RGBColor(255, 255, 255),
            "title": RGBColor(31, 49, 74),
            "text": RGBColor(38, 45, 56),
            "accent": RGBColor(53, 100, 160),
            "muted": RGBColor(98, 108, 120),
        }
    if theme == "minimal":
        return {
            "bg": RGBColor(255, 255, 255),
            "panel": RGBColor(255, 255, 255),
            "title": RGBColor(20, 20, 22),
            "text": RGBColor(45, 45, 50),
            "accent": RGBColor(90, 90, 100),
            "muted": RGBColor(125, 125, 135),
        }
    return {
        "bg": RGBColor(24, 25, 31),
        "panel": RGBColor(34, 36, 44),
        "title": RGBColor(248, 248, 250),
        "text": RGBColor(232, 233, 238),
        "accent": RGBColor(215, 97, 150),
        "muted": RGBColor(170, 173, 184),
    }


def _set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text(slide, text: str, x: float, y: float, w: float, h: float, *, size: int,
              color: RGBColor, bold: bool = False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Microsoft JhengHei"
    return box


def _add_bullets(slide, bullets: list[str], colors: dict[str, RGBColor]) -> None:
    if not bullets:
        _add_text(
            slide,
            "重點內容請搭配口說稿說明。",
            1.05, 2.05, 11.15, 3.8,
            size=24,
            color=colors["muted"],
        )
        return
    box = slide.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.2), Inches(4.8))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    font_size = 26 if len(bullets) <= 5 else 22
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.space_after = Pt(12)
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = colors["text"]
        paragraph.font.name = "Microsoft JhengHei"
        paragraph.text = f"• {bullet}"


def _build_pptx(title: str, subtitle: str, slides_data: list[dict[str, object]], theme_name: str, path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    colors = _theme(theme_name)

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _set_bg(title_slide, colors["bg"])
    accent = title_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.2), Inches(0.12), Inches(4.7))
    accent.fill.solid()
    accent.fill.fore_color.rgb = colors["accent"]
    accent.line.fill.background()
    _add_text(title_slide, title, 1.15, 1.45, 11.2, 1.8, size=34, color=colors["title"], bold=True)
    if subtitle:
        _add_text(title_slide, subtitle, 1.18, 3.35, 10.8, 1.1, size=20, color=colors["muted"])
    _add_text(title_slide, "MOXUE PRESENTATION", 1.18, 5.55, 5.0, 0.45, size=11, color=colors["accent"], bold=True)

    total = len(slides_data)
    for index, item in enumerate(slides_data, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _set_bg(slide, colors["bg"])
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Inches(0.14))
        header.fill.solid()
        header.fill.fore_color.rgb = colors["accent"]
        header.line.fill.background()
        _add_text(slide, str(item["title"]), 0.95, 0.55, 11.3, 0.9, size=27, color=colors["title"], bold=True)
        _add_bullets(slide, list(item["bullets"]), colors)
        _add_text(slide, f"{index:02d} / {total:02d}", 11.55, 6.95, 1.0, 0.3, size=10, color=colors["muted"], align=PP_ALIGN.RIGHT)

    presentation.save(path)


def _build_script_docx(title: str, subtitle: str, slides_data: list[dict[str, object]], path: Path) -> None:
    document = Document()
    document.core_properties.title = f"{title}－口說稿"
    document.add_heading(title, level=0)
    if subtitle:
        document.add_paragraph(subtitle)
    document.add_paragraph("以下依投影片順序整理口說重點與講稿。")
    for index, item in enumerate(slides_data, start=1):
        document.add_heading(f"Slide {index}｜{item['title']}", level=1)
        bullets = list(item["bullets"])
        if bullets:
            document.add_heading("投影片重點", level=2)
            for bullet in bullets:
                document.add_paragraph(str(bullet), style="List Bullet")
        document.add_heading("口說稿", level=2)
        notes = str(item["speaker_notes"]).strip()
        if not notes:
            notes = "；".join(str(value) for value in bullets) or "依投影片內容自然說明。"
        for paragraph in notes.split("\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
    styles = document.styles
    styles["Normal"].font.name = "Microsoft JhengHei"
    styles["Normal"].font.size = DocxPt(11)
    document.save(path)


def _sections(arguments: dict[str, object]) -> list[dict[str, object]]:
    raw = arguments.get("sections")
    if not isinstance(raw, list) or not raw:
        raise ValueError("sections must be a non-empty array.")
    result = []
    for index, item in enumerate(raw[:30], start=1):
        if not isinstance(item, dict):
            raise ValueError(f"sections[{index}] must be an object.")
        heading = _text(item.get("heading"), label=f"sections[{index}].heading", maximum=180)
        body = _text(item.get("body", ""), label=f"sections[{index}].body", maximum=MAX_SECTION_BODY, allow_empty=True)
        bullets_raw = item.get("bullets", [])
        if not isinstance(bullets_raw, list):
            raise ValueError(f"sections[{index}].bullets must be an array.")
        bullets = [
            _text(value, label=f"sections[{index}].bullets", maximum=MAX_BULLET_CHARS)
            for value in bullets_raw[:MAX_BULLETS]
            if str(value or "").strip()
        ]
        result.append({"heading": heading, "body": body, "bullets": bullets})
    return result


def _build_document(title: str, sections: list[dict[str, object]], docx_path: Path, md_path: Path) -> None:
    document = Document()
    document.core_properties.title = title
    document.add_heading(title, level=0)
    markdown = [f"# {title}", ""]
    for section in sections:
        heading = str(section["heading"])
        body = str(section["body"])
        bullets = list(section["bullets"])
        document.add_heading(heading, level=1)
        markdown.extend([f"## {heading}", ""])
        if body:
            for paragraph in body.split("\n"):
                if paragraph.strip():
                    document.add_paragraph(paragraph.strip())
                    markdown.extend([paragraph.strip(), ""])
        for bullet in bullets:
            document.add_paragraph(str(bullet), style="List Bullet")
            markdown.append(f"- {bullet}")
        if bullets:
            markdown.append("")
    document.styles["Normal"].font.name = "Microsoft JhengHei"
    document.styles["Normal"].font.size = DocxPt(11)
    document.save(docx_path)
    md_path.write_text("\n".join(markdown).rstrip() + "\n", encoding="utf-8")


async def _create_presentation_package(arguments: dict[str, object]) -> dict[str, object]:
    title = _text(arguments.get("title"), label="title", maximum=MAX_TITLE)
    subtitle = _text(arguments.get("subtitle", ""), label="subtitle", maximum=500, allow_empty=True)
    slides_data = _slides(arguments)
    theme_name = str(arguments.get("theme", "modern")).strip().lower() or "modern"
    if theme_name not in {"modern", "academic", "minimal"}:
        raise ValueError("theme must be modern, academic, or minimal.")
    stem = _safe_stem(arguments.get("filename_stem"), _safe_stem(title, "presentation"))
    token = uuid.uuid4().hex[:10]
    directory = _output_dir()
    pptx_path = directory / f"{stem}-{token}.pptx"
    docx_path = directory / f"{stem}-speaker-script-{token}.docx"
    _build_pptx(title, subtitle, slides_data, theme_name, pptx_path)
    _build_script_docx(title, subtitle, slides_data, docx_path)
    return {
        "message": "簡報與口說稿已建立，請用核心附件機制交付給使用者。",
        "_moxue_artifacts": [
            _artifact(pptx_path, filename=f"{stem}.pptx", mime_type=PPTX_MIME),
            _artifact(docx_path, filename=f"{stem}_口說稿.docx", mime_type=DOCX_MIME),
        ],
    }


async def _create_document(arguments: dict[str, object]) -> dict[str, object]:
    title = _text(arguments.get("title"), label="title", maximum=MAX_TITLE)
    sections = _sections(arguments)
    stem = _safe_stem(arguments.get("filename_stem"), _safe_stem(title, "document"))
    token = uuid.uuid4().hex[:10]
    directory = _output_dir()
    docx_path = directory / f"{stem}-{token}.docx"
    md_path = directory / f"{stem}-{token}.md"
    _build_document(title, sections, docx_path, md_path)
    return {
        "message": "Word 與 Markdown 文件已建立，請用核心附件機制交付給使用者。",
        "_moxue_artifacts": [
            _artifact(docx_path, filename=f"{stem}.docx", mime_type=DOCX_MIME),
            _artifact(md_path, filename=f"{stem}.md", mime_type=MARKDOWN_MIME),
        ],
    }


async def invoke(action: str, arguments: dict[str, object], context: dict[str, object]) -> object:
    del context
    try:
        if action == "create_presentation_package":
            return await _create_presentation_package(arguments)
        if action == "create_document":
            return await _create_document(arguments)
        raise ValueError(f"Unknown action: {action}")
    except (ValueError, DocumentGenerationError) as error:
        return {"message": str(error), "error": True}
