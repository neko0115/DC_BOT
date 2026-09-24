from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from pptx import Presentation

from discord_ai_assistant.meeting_report_exports import build_meeting_report_exports


SAMPLE_REPORT = """> 📅 **週會報日期：** 2026-09-24
> 🕒 **產生時間：** 14:30 (CST)

# 明日之後週會

## 本週進度

- 完成週會錄音整理
- 修正專有名詞辨識

## 問題與阻塞

1. 尚未完成實機驗證
2. 下週需要再次測試

## 下一步

持續確認 Discord Bot 的週會輸出。
"""


class MeetingReportExportTests(unittest.TestCase):
    def test_builds_markdown_docx_and_pptx_from_same_report(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            report = root / "report_20260924.md"
            report.write_text(SAMPLE_REPORT, encoding="utf-8")

            exports = build_meeting_report_exports(report, title="明日之後週會")

            self.assertEqual(exports.markdown, report)
            self.assertTrue(exports.docx.exists())
            self.assertTrue(exports.pptx.exists())
            self.assertGreater(exports.docx.stat().st_size, 1_000)
            self.assertGreater(exports.pptx.stat().st_size, 1_000)

            document = Document(exports.docx)
            docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("明日之後週會", docx_text)
            self.assertIn("完成週會錄音整理", docx_text)
            self.assertIn("尚未完成實機驗證", docx_text)

            presentation = Presentation(exports.pptx)
            slide_text = "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if hasattr(shape, "text")
            )
            self.assertIn("明日之後週會", slide_text)
            self.assertIn("本週進度", slide_text)
            self.assertIn("完成週會錄音整理", slide_text)
            self.assertIn("下一步", slide_text)

    def test_empty_report_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            report = Path(directory) / "empty.md"
            report.write_text(" \n", encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "週會報內容為空"):
                build_meeting_report_exports(report, title="空白週會")


if __name__ == "__main__":
    unittest.main()
