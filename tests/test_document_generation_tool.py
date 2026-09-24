from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
TOOL_PATH = ROOT / "tools" / "document_generation" / "tool.py"


def load_tool():
    spec = importlib.util.spec_from_file_location("moxue_document_generation_tool", TOOL_PATH)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class DocumentGenerationToolTests(unittest.IsolatedAsyncioTestCase):
    async def test_presentation_package_builds_openable_pptx_and_speaker_script(self) -> None:
        tool = load_tool()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch.object(tool, "_output_dir", return_value=root):
                result = await tool.invoke(
                    "create_presentation_package",
                    {
                        "title": "英式藍莓煎餅",
                        "subtitle": "課堂料理介紹",
                        "theme": "modern",
                        "slides": [
                            {
                                "title": "什麼是英式藍莓煎餅",
                                "bullets": ["以麵糊與藍莓製作", "適合作為早餐或點心"],
                                "speaker_notes": "先介紹料理定位，再說明今天的報告架構。",
                            },
                            {
                                "title": "製作步驟",
                                "bullets": ["混合材料", "加入藍莓", "煎至兩面上色"],
                                "speaker_notes": "依序說明三個步驟與火候注意事項。",
                            },
                        ],
                    },
                    {},
                )

            self.assertFalse(result.get("error", False))
            artifacts = result["_moxue_artifacts"]
            self.assertEqual(len(artifacts), 2)
            pptx = root / Path(artifacts[0]["relative_path"]).name
            docx = root / Path(artifacts[1]["relative_path"]).name
            self.assertTrue(pptx.exists())
            self.assertTrue(docx.exists())

            presentation = Presentation(pptx)
            slide_text = "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if hasattr(shape, "text")
            )
            self.assertIn("英式藍莓煎餅", slide_text)
            self.assertIn("製作步驟", slide_text)

            document = Document(docx)
            doc_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("口說稿", doc_text)
            self.assertIn("火候注意事項", doc_text)

    async def test_document_builds_matching_docx_and_markdown(self) -> None:
        tool = load_tool()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch.object(tool, "_output_dir", return_value=root):
                result = await tool.invoke(
                    "create_document",
                    {
                        "title": "測試文件",
                        "sections": [
                            {
                                "heading": "摘要",
                                "body": "這是一份測試文件。",
                                "bullets": ["重點一", "重點二"],
                            }
                        ],
                    },
                    {},
                )

            artifacts = result["_moxue_artifacts"]
            docx = root / Path(artifacts[0]["relative_path"]).name
            md = root / Path(artifacts[1]["relative_path"]).name
            self.assertIn("測試文件", "\n".join(p.text for p in Document(docx).paragraphs))
            self.assertIn("## 摘要", md.read_text(encoding="utf-8"))
            self.assertIn("- 重點二", md.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
