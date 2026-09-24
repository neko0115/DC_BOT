from __future__ import annotations

import importlib.util
import os
import tempfile
import unittest
from pathlib import Path

from docx import Document
from pptx import Presentation

from discord_ai_assistant.ai.tools import ToolRouter

ROOT = Path(__file__).resolve().parents[1]
DELIVERY_PATH = ROOT / "src" / "discord_ai_assistant" / "artifact_delivery.py"


def load_delivery():
    if not DELIVERY_PATH.is_file():
        raise AssertionError("artifact_delivery.py has not been implemented")
    spec = importlib.util.spec_from_file_location("moxue_artifact_delivery", DELIVERY_PATH)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class ArtifactDeliveryContractTests(unittest.TestCase):
    def test_tool_router_hides_artifact_metadata_and_translates_to_core_effect(self) -> None:
        result = ToolRouter._attach_external_effects(
            {
                "message": "done",
                "_moxue_artifacts": [
                    {
                        "relative_path": "image_generation/abc.jpg",
                        "filename": "generated.jpg",
                        "mime_type": "image/jpeg",
                        "delete_after_send": True,
                    }
                ],
            }
        )
        self.assertNotIn("_moxue_artifacts", result)
        self.assertEqual(
            result["_moxue_effects"],
            [
                {
                    "type": "attach_artifact",
                    "relative_path": "image_generation/abc.jpg",
                    "filename": "generated.jpg",
                    "mime_type": "image/jpeg",
                    "delete_after_send": True,
                }
            ],
        )

    def test_plugin_cannot_inject_core_effects_directly(self) -> None:
        result = ToolRouter._attach_external_effects(
            {
                "message": "attempt",
                "_moxue_effects": [
                    {"type": "publish_final_reply", "channel_id": 999, "suppress_origin": True}
                ],
            }
        )
        self.assertNotIn("_moxue_effects", result)

    def test_router_rejects_artifact_path_traversal(self) -> None:
        result = ToolRouter._attach_external_effects(
            {
                "_moxue_artifacts": [
                    {
                        "relative_path": "../secrets.env",
                        "filename": "secrets.env",
                        "mime_type": "text/plain",
                    }
                ]
            }
        )
        self.assertNotIn("_moxue_effects", result)

    def test_core_resolver_accepts_valid_office_and_markdown_artifacts(self) -> None:
        delivery = load_delivery()
        with tempfile.TemporaryDirectory() as directory:
            project_root = Path(directory)
            root = project_root / "data" / "tool-artifacts" / "document_generation"
            root.mkdir(parents=True)

            docx_path = root / "report.docx"
            document = Document()
            document.add_paragraph("測試 Word")
            document.save(docx_path)

            pptx_path = root / "slides.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = "測試簡報"
            presentation.save(pptx_path)

            md_path = root / "notes.md"
            md_path.write_text("# 測試 Markdown\n", encoding="utf-8")

            effects = [
                {
                    "type": "attach_artifact",
                    "relative_path": "document_generation/report.docx",
                    "filename": "report.docx",
                    "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                },
                {
                    "type": "attach_artifact",
                    "relative_path": "document_generation/slides.pptx",
                    "filename": "slides.pptx",
                    "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                },
                {
                    "type": "attach_artifact",
                    "relative_path": "document_generation/notes.md",
                    "filename": "notes.md",
                    "mime_type": "text/markdown",
                },
            ]
            for effect in effects:
                self.assertIsNotNone(
                    delivery.resolve_artifact_effect(project_root, 25 * 1024 * 1024, effect)
                )

            fake_docx = root / "fake.docx"
            fake_docx.write_bytes(b"PK but not really an Office archive")
            fake_effect = dict(effects[0], relative_path="document_generation/fake.docx")
            self.assertIsNone(
                delivery.resolve_artifact_effect(project_root, 25 * 1024 * 1024, fake_effect)
            )

    def test_core_resolver_enforces_root_mime_size_magic_and_symlink(self) -> None:
        delivery = load_delivery()
        with tempfile.TemporaryDirectory() as directory:
            project_root = Path(directory)
            root = project_root / "data" / "tool-artifacts" / "image_generation"
            root.mkdir(parents=True)
            valid = root / "valid.jpg"
            valid.write_bytes(b"\xff\xd8\xffpayload")

            effect = {
                "type": "attach_artifact",
                "relative_path": "image_generation/valid.jpg",
                "filename": "generated.jpg",
                "mime_type": "image/jpeg",
                "delete_after_send": True,
            }
            resolved = delivery.resolve_artifact_effect(project_root, 25 * 1024 * 1024, effect)
            self.assertIsNotNone(resolved)

            traversal = dict(effect, relative_path="../valid.jpg")
            self.assertIsNone(delivery.resolve_artifact_effect(project_root, 25 * 1024 * 1024, traversal))

            mismatch = dict(effect, mime_type="image/png")
            self.assertIsNone(delivery.resolve_artifact_effect(project_root, 25 * 1024 * 1024, mismatch))

            fake_jpeg = root / "fake.jpg"
            fake_jpeg.write_bytes(b"this is not a jpeg")
            fake_effect = dict(effect, relative_path="image_generation/fake.jpg")
            self.assertIsNone(delivery.resolve_artifact_effect(project_root, 25 * 1024 * 1024, fake_effect))

            self.assertIsNone(delivery.resolve_artifact_effect(project_root, 4, effect))

            link = root / "link.jpg"
            try:
                os.symlink(valid, link)
            except (OSError, NotImplementedError):
                return
            link_effect = dict(effect, relative_path="image_generation/link.jpg")
            self.assertIsNone(delivery.resolve_artifact_effect(project_root, 25 * 1024 * 1024, link_effect))


if __name__ == "__main__":
    unittest.main()
